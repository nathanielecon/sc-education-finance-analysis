from __future__ import annotations

import csv
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


def read_pdf_text(path: Path) -> str:
    """Read text from a generic PDF input."""
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def read_html_rows(path: Path) -> list[list[str]]:
    """Read rows from the first table in a generic HTML input."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    table = soup.find("table")
    if table is None:
        raise ValueError("The HTML input does not contain a table.")
    return [
        [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
        for row in table.find_all("tr")
    ]


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    """Read a generic CSV input into named rows."""
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def read_workbook_rows(path: Path) -> list[list[object]]:
    """Read values from the active worksheet in a generic workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    return [list(row) for row in workbook.active.iter_rows(values_only=True)]
